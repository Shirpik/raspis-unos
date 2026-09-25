import glob
import os
import re
import json
import datetime as dt
from collections import Counter, defaultdict

import openpyxl

TEMPLATE = r"C:\Users\Student\Downloads\_Учёт часов в СПО 2 семестр (1).xlsx"
SOURCE_DIR = "data"
CUTOFF = dt.date(2026, 9, 22)


def norm(value):
    return re.sub(r"[^а-яё]", "", str(value).lower())


def parse_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        value = value.strip()
        for fmt in ("%d.%m.%Y", "%d.%m"):
            try:
                parsed = dt.datetime.strptime(value, fmt)
                return parsed.date().replace(year=2026)
            except ValueError:
                pass
    return None


WEEKDAYS = {
    "понедельник": 0, "вт": 1, "вторник": 1, "ср": 2, "среда": 2,
    "четверг": 3, "чт": 3, "пятница": 4, "суббота": 5,
}


def correct_weekday_date(sheet, row, candidate):
    """Fix a stale date formula by reconciling the displayed day of week."""
    label = str(sheet.cell(row + 1, 1).value or "").strip().lower()
    target = WEEKDAYS.get(label)
    if target is None:
        return candidate
    for shift in (0, 1, -1, 2, -2, 3, -3):
        possible = candidate + dt.timedelta(days=shift)
        if possible.weekday() == target:
            return possible
    return candidate


def teacher_database():
    workbook = openpyxl.load_workbook(TEMPLATE, data_only=True, read_only=True)
    names = set()
    for sheet in workbook.worksheets[2:]:
        for row in sheet.iter_rows(min_col=5, max_col=5, values_only=True):
            value = row[0]
            if isinstance(value, str) and len(value.split()) >= 2:
                names.add(value.strip())
    return sorted(names)


def teacher_plans():
    workbook = openpyxl.load_workbook(TEMPLATE, data_only=True, read_only=True)
    plans = defaultdict(float)
    for sheet in workbook.worksheets[2:]:
        for row in sheet.iter_rows(min_col=4, max_col=5, values_only=True):
            hours, teacher = row
            if isinstance(hours, (int, float)) and isinstance(teacher, str) and len(teacher.split()) >= 2:
                plans[teacher.strip()] += hours
    return dict(plans)


def find_teacher(cell, teachers):
    normalized = norm(cell)
    matches = [name for name in teachers if norm(name) and norm(name) in normalized]
    if len(matches) == 1:
        return matches[0]
    # Schedules normally show surname only; use it when the database has one such surname.
    surname_matches = [name for name in teachers if norm(name.split()[0]) in normalized]
    if len(surname_matches) == 1:
        return surname_matches[0]
    return None


def get_subject(cell, teacher):
    lines = [x.strip() for x in str(cell).split("\n") if x.strip()]
    teacher_line = next((i for i, line in enumerate(lines) if norm(teacher) in norm(line)), None)
    if teacher_line is None:
        return lines[0] if lines else ""
    return " ".join(lines[:teacher_line]).strip() or lines[0]


def collect():
    teachers = teacher_database()
    records = []
    unmatched = Counter()
    for path in sorted(glob.glob(os.path.join(SOURCE_DIR, "*.xlsx"))):
        source = os.path.basename(path).removesuffix(".xlsx")
        book = openpyxl.load_workbook(path, data_only=True)
        for sheet in book.worksheets:
            groups = {column: sheet.cell(1, column).value for column in range(4, sheet.max_column + 1)}
            current_date = parse_date(sheet.cell(1, 1).value)
            for row in range(1, sheet.max_row + 1):
                candidate = parse_date(sheet.cell(row, 1).value)
                if candidate:
                    current_date = correct_weekday_date(sheet, row, candidate)
                    for column in range(4, sheet.max_column + 1):
                        group = sheet.cell(row, column).value
                        if isinstance(group, str) and group.strip():
                            groups[column] = group.strip()
                    continue
                label = str(sheet.cell(row, 1).value or "").strip().lower()
                target_weekday = WEEKDAYS.get(label)
                if current_date and target_weekday is not None and current_date.weekday() != target_weekday:
                    for shift in range(1, 7):
                        possible = current_date + dt.timedelta(days=shift)
                        if possible.weekday() == target_weekday:
                            current_date = possible
                            break
                if not current_date or current_date > CUTOFF:
                    continue
                for column, group in groups.items():
                    value = sheet.cell(row, column).value
                    if not isinstance(value, str) or not value.strip():
                        continue
                    if value.strip() in {"#REF!", "#VALUE!"}:
                        continue
                    teacher = find_teacher(value, teachers)
                    if teacher:
                        records.append({
                            "date": current_date.isoformat(),
                            "course": sheet.title,
                            "group": str(group).strip(),
                            "teacher": teacher,
                            "subject": get_subject(value, teacher),
                            "hours": 2,
                            "source": source,
                        })
                    elif "\n" in value:
                        unmatched[value.split("\n")[-1].strip()] += 1
    return teachers, records, unmatched


if __name__ == "__main__":
    teachers, records, unmatched = collect()
    plans = teacher_plans()
    summary = defaultdict(int)
    groups = defaultdict(int)
    for item in records:
        summary[item["teacher"]] += item["hours"]
        groups[(item["teacher"], item["group"])] += item["hours"]
    report = {
        "teacher_count": len(teachers),
        "matched_events": len(records),
        "matched_hours": sum(item["hours"] for item in records),
        "active_teachers": len(summary),
        "plans": plans,
        "top_teachers": sorted(summary.items(), key=lambda x: (-x[1], x[0]))[:20],
        "unmatched_top": unmatched.most_common(30),
        "records": records,
    }
    with open("analysis.json", "w", encoding="utf-8") as output:
        json.dump(report, output, ensure_ascii=False, indent=2)
    print(json.dumps({k: v for k, v in report.items() if k != "records"}, ensure_ascii=False, indent=2))
