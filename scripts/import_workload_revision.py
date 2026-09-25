"""Prepare and verify a stable-ID semester workload import before atomic install.

Workbook is read-only. Removed/zero-semester rows are archived, never deleted;
history, calendars, availability and existing lesson settings are retained.
"""
import argparse, copy, datetime as dt, hashlib, json, os, subprocess
from collections import defaultdict
from pathlib import Path
import openpyxl
from merge_vkleyki_workload import clean, norm, numeric, columns, subgroup, subject_name

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT/'data/timetable_data.json'
def digest(content): return hashlib.sha256(content).hexdigest()
def write(path, value): path.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def identity(row): return row['group'], norm(row['name']), row.get('subgroup',-1)
def read_rows(source, data):
    groups={norm(g['name']):g for g in data['groups']}
    wb=openpyxl.load_workbook(source,data_only=True)
    assert {norm(s) for s in wb.sheetnames}==set(groups), 'Workbook must cover all known groups'
    rows=[]; all_rows=[]
    for ws in wb:
        group=groups[norm(ws.title)]; h,s,t,c=columns(ws)
        assert min(h,s,t)>0, ws.title
        assert any(col==c and any(str(n) in clean(ws.cell(r,col).value) for n in (1,3,5,7))
                   for r in range(1,13) for col in range(1,ws.max_column+1)
                   if 'сем' in norm(ws.cell(r,col).value)), ('No odd-semester column',ws.title)
        for r in range(h+1,ws.max_row+1):
            raw=clean(ws.cell(r,s).value)
            if not raw: continue
            hours=numeric(ws.cell(r,c).value)
            row=dict(group=group['id'],name=subject_name(raw),subgroup=subgroup(raw,group['id']),
                     total_hours=hours or 0,teacher_name=clean(ws.cell(r,t).value),
                     source_index=clean(ws.cell(r,2).value),sheet=ws.title,row=r)
            all_rows.append(row)
            if hours is not None and hours>0:
                assert hours==int(hours), ('Fractional hours',row)
                row['total_hours']=int(hours); rows.append(row)
    assert len({identity(r) for r in rows})==len(rows), 'Duplicate curriculum identities'
    return rows,all_rows

def prepare(source,out):
    raw=DATA.read_bytes(); before=json.loads(raw.decode('utf-8-sig')); data=copy.deepcopy(before)
    rows,all_rows=read_rows(source,data)
    teachers={norm(t['name']):t['id'] for t in data['teachers']}
    old_by_key=defaultdict(list)
    for l in data['lessons']: old_by_key[identity(l)].append(l)
    used=set(); changes=[]; new=[]; archived=[]; result=[]
    next_id=max(l['id'] for l in data['lessons'])+1
    next_subject=max(l.get('subject_id',-1) for l in data['lessons'])+1
    stamp=dt.datetime.now(dt.timezone.utc).isoformat()
    for row in rows:
        teacher=row['teacher_name']; tkey=norm(teacher)
        if not teacher or 'вакансия' in tkey or tkey=='вынесена на пп': tid=-1
        else:
            assert tkey in teachers, ('Unknown teacher requires reconciliation',row)
            tid=teachers[tkey]
        candidates=[l for l in old_by_key[identity(row)] if l['id'] not in used]
        exact=[l for l in candidates if l.get('teacher',-1)==tid]
        candidates=exact or candidates
        if not candidates and row['source_index']:
            # A name change can retain its ID only when group/index/subgroup/kind
            # identify exactly one old row, which is absent under its old name.
            candidates=[l for l in data['lessons'] if l['id'] not in used and
                        l['group']==row['group'] and l.get('subgroup',-1)==row['subgroup'] and
                        norm(l.get('source_index'))==norm(row['source_index']) and
                        bool(l.get('is_lab'))==('лпз' in norm(row['name'])) and
                        not any(identity(l)==identity(r) for r in rows)]
        assert len(candidates)<=1, ('Ambiguous identity',row,candidates)
        if candidates:
            previous=candidates[0]; item=copy.deepcopy(previous); used.add(item['id'])
        else:
            previous=None
            item=dict(id=next_id,uid='lesson-'+digest(f"{source.name}|{identity(row)}".encode())[:16],
                      group=row['group'],subgroup=row['subgroup'],subject_id=next_subject,
                      total_slots=0,generation_active=False,curriculum_active=True,plan_active=True,
                      is_lab='лпз' in norm(row['name']),is_block=False,is_pp=False,
                      allowed_campuses=[0,1],week_parity='all',fixed_room=-1,preferred_room=-1,
                      allow_room_substitution=True,required_room_type=0,required_capacity=0,
                      required_equipment=[],consecutive_pairs=1)
            # New curriculum rows are accepted from the source as-is.  Their
            # source index/subgroup identity is retained for the next audit.
            next_id+=1; next_subject+=1
        values=dict(name=row['name'],teacher=tid,total_hours=row['total_hours'],source_index=row['source_index'])
        diff={k:[previous.get(k),v] for k,v in values.items() if previous and previous.get(k)!=v}
        item.update(values)
        if previous and previous.get('teacher',-1)!=tid:
            # A quota selected for the former assignee is not a valid quota for
            # their replacement. Ledger/saved schedules keep actual teachers.
            for field,value in [('total_slots',0),('generation_active',False)]:
                if item.get(field)!=value:
                    diff[field]=[item.get(field),value]
                    item[field]=value
        item['curriculum_active']=True
        item['workload_source']={'file':source.name,'sheet':row['sheet'],'row':row['row'],'semester':1}
        if previous and diff: changes.append(dict(id=item['id'],group=row['sheet'],name=item['name'],changes=diff))
        if previous is None: new.append(dict(id=item['id'],group=row['sheet'],name=item['name'],hours=item['total_hours']))
        result.append(item)
    for previous in data['lessons']:
        if previous['id'] in used: continue
        item=copy.deepcopy(previous)
        source_rows=[r for r in all_rows if identity(r)==identity(previous)]
        replacement=[r for r in rows if r['group']==previous['group'] and r['subgroup']==previous.get('subgroup',-1)
                     and norm(r['source_index'])==norm(previous.get('source_index')) and r['source_index']]
        assert (source_rows and all(r['total_hours']==0 for r in source_rows)) or len(replacement)==1, ('Unexplained removed row',previous)
        reason='В первом семестре часов нет' if source_rows else 'Заменена предметом из новых вклеек: '+replacement[0]['name']
        item.update(total_hours=0,total_slots=0,curriculum_active=False,plan_active=False,generation_active=False)
        item['workload_retired']={'file':source.name,'at':stamp,'previous_total_hours':previous.get('total_hours',0),'reason':reason}
        archived.append(dict(id=item['id'],name=item['name'],group=next(g['name'] for g in data['groups'] if g['id']==item['group']),reason=reason,old_hours=previous.get('total_hours',0)))
        result.append(item)
    data['lessons']=sorted(result,key=lambda l:l['id'])
    assert len({l['id'] for l in result})==len(result)
    assert {l['id'] for l in before['lessons']}<={l['id'] for l in result}
    # Existing historical teacher snapshots must survive assignment transfers.
    transferred={c['id'] for c in changes if 'teacher' in c['changes']}
    assert all('actual_teacher' in r for r in before['teaching_ledger'] if r['lesson_id'] in transferred)
    metadata=dict(id=max((x.get('id',-1) for x in data.get('workload_imports',[])),default=-1)+1,
                  file_name=source.name,semester=1,imported_at=stamp,groups=len(data['groups']),
                  teachers=len(data['teachers']),active_lessons=len(rows),mode='reconciled_stable_id',sha256=digest(source.read_bytes()))
    data.setdefault('workload_imports',[]).append(metadata)
    for k in before:
        if k not in ('lessons','workload_imports'): assert before[k]==data[k], k
    report=dict(source=str(source),source_sha256=metadata['sha256'],baseline_sha256=digest(raw),
                groups=len(data['groups']),source_rows=len(rows),hours=sum(r['total_hours'] for r in rows),
                changes=changes,new=new,archived=archived,
                ledger_records_preserved=len(data['teaching_ledger']),teacher_constraints_preserved=True)
    out.mkdir(parents=True,exist_ok=True)
    (out/'before-data.json').write_bytes(raw)
    write(out/'candidate.json',data); write(out/'import-report.json',report)
    exe=ROOT/'.tmp/build-sep7/Release/timetable_solver.exe'
    for label,file in [('before',out/'before-data.json'),('after',out/'candidate.json')]:
        for task in ('audit-data','semester-readout'):
            run=subprocess.run([str(exe),'--'+task,'--data',str(file)],capture_output=True)
            (out/f'{label}-{task}.json').write_bytes(run.stdout)
            assert run.returncode in (0,1) and run.stdout.strip().startswith(b'{'), run.stderr.decode('utf-8',errors='replace')
    after_audit=json.loads((out/'after-audit-data.json').read_bytes())
    before_audit=json.loads((out/'before-audit-data.json').read_bytes())
    report['audit_before']=before_audit.get('summary')
    report['audit_after']=after_audit.get('summary')
    report['candidate_sha256']=digest((out/'candidate.json').read_bytes())
    write(out/'import-report.json',report)
    print(json.dumps(report,ensure_ascii=False,indent=2))

def commit(source,out):
    report=json.loads((out/'import-report.json').read_bytes())
    assert report['audit_after']['errors']==0, 'Resolve structural/selected-quota errors before install'
    assert digest(DATA.read_bytes())==report['baseline_sha256'], 'Database changed since preparation'
    assert digest(source.read_bytes())==report['source_sha256'], 'Source changed since preparation'
    candidate=(out/'candidate.json').read_bytes()
    assert digest(candidate)==report['candidate_sha256'], 'Candidate changed since audit'
    archive=ROOT/'data/history'/('before_vkleyki_v5_'+dt.datetime.now().strftime('%Y%m%d_%H%M%S')+'.json')
    archive.parent.mkdir(exist_ok=True)
    archive.write_bytes(DATA.read_bytes())
    temporary=DATA.with_suffix('.json.workload-import.tmp')
    temporary.write_bytes(candidate)
    os.replace(temporary,DATA)
    assert digest(DATA.read_bytes())==report['candidate_sha256']
    write(out/'installed.json',dict(installed_at=dt.datetime.now(dt.timezone.utc).isoformat(),backup=str(archive),sha256=report['candidate_sha256']))
    print('Installed:',str(DATA))

if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('source',type=Path); parser.add_argument('--out',type=Path,default=ROOT/'outputs/workload-import-20260914-v5')
    parser.add_argument('--commit',action='store_true')
    args=parser.parse_args()
    (commit if args.commit else prepare)(args.source.resolve(),args.out.resolve())
