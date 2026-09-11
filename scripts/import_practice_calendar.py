"""Read the supplied annual calendar. Never modify the source workbook.

Dates come from the full-date header on sheet one, checked against day numbers
on sheet two. Practice is identified by its named row, not by green alone:
the same workbook also uses green for holidays and theme colours for PP.
"""
import argparse
import copy
import json
import re
import hashlib
from datetime import date, timedelta
from pathlib import Path

import openpyxl


def key(value):
    return re.sub(r'[^а-яa-z0-9]', '', str(value).lower().replace('ё', 'е')).replace('пф', '')


ALIASES = {'мцм102': 'МЦМ-Пф-103', 'тэорп2901': 'ТЭиРП-2901', 'такхс2202': 'ТАКХС-Пф-2202',
           'пкд380607п': 'ПКД-380607п', 'такх4202': 'ТАКХС-4202'}
CONFIRMED_ALIASES = {'такхс3202': 'ТАКХС-3201', 'мцм407': 'МЦМ-408'}


def extract(source, data, confirm_aliases=False):
    workbook = openpyxl.load_workbook(source, data_only=True)
    headers = workbook['1_2_курсы']
    weeks = []
    for col in range(9, 61):
        first, last = headers.cell(10, col).value, headers.cell(11, col).value
        if not hasattr(first, 'date') or not hasattr(last, 'date'):
            raise ValueError(f'Не прочитаны полные даты столбца {col}')
        weeks.append((col, first.date(), last.date()))
        if last.date() - first.date() != timedelta(days=6) or first.weekday() != 0:
            raise ValueError(f'Некорректная учебная неделя в столбце {col}')
        if len(weeks) > 1 and weeks[-2][2] + timedelta(days=1) != first.date():
            raise ValueError(f'Разрыв календаря в столбце {col}')
    groups = {key(g['name']): g for g in data['groups']}
    aliases = {**ALIASES, **(CONFIRMED_ALIASES if confirm_aliases else {})}
    entries, unmatched, duplicates = [], [], []
    used = set()
    for sheet in workbook:
        if sheet.title == '3_4_курсы':
            for col, first, last in weeks:
                if int(sheet.cell(2, col).value) != first.day or int(sheet.cell(3, col).value) != last.day:
                    raise ValueError(f'Даты листов расходятся в столбце {col}')
        starts = [r for r in range(1, sheet.max_row + 1)
                  if sheet.cell(r, 3).value and str(sheet.cell(r, 5).value).strip() == 'теория']
        for pos, row in enumerate(starts):
            name = str(sheet.cell(row, 3).value).strip()
            normalized = key(name)
            mapped = groups.get(key(aliases.get(normalized, name)))
            if mapped and mapped['id'] in used:
                # Two similarly named SP groups are different entities. Never
                # silently rename a second occurrence based on row order.
                duplicates.append({'sheet': sheet.title, 'row': row, 'name': name})
                continue
            if not mapped:
                unmatched.append({'sheet': sheet.title, 'row': row, 'name': name})
                continue
            used.add(mapped['id'])
            stop = starts[pos + 1] if pos + 1 < len(starts) else min(row + 8, sheet.max_row + 1)
            practice_rows = [r for r in range(row + 1, stop)
                             if str(sheet.cell(r, 5).value).strip().replace(' ', '') in ('п.практика', 'пр.практика', 'ПДП')]
            practice_weeks = []
            calendar_weeks = []
            for col, first, last in weeks:
                record = {'from': first.isoformat(), 'to': last.isoformat(),
                          'theory_hours': 0, 'up_hours': 0, 'pp_hours': 0,
                          'exam_hours': 0, 'vacation': False, 'source_cells': []}
                for r in range(row, stop):
                    cell = sheet.cell(r, col)
                    label = str(sheet.cell(r, 5).value).strip().lower().replace(' ', '')
                    field = {'теория': 'theory_hours', 'уч.практика': 'up_hours',
                             'п.практика': 'pp_hours', 'пр.практика': 'pp_hours',
                             'пдп': 'pp_hours', 'экзамены': 'exam_hours', 'гиа': 'exam_hours'}.get(label)
                    if not field:
                        continue
                    value = cell.value
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        if value < 0 or value != int(value):
                            raise ValueError(f'{sheet.title}!{cell.coordinate}: некорректные часы')
                        record[field] += int(value)
                    if label == 'теория' and str(value).strip() == '=':
                        record['vacation'] = True
                    if value is not None:
                        record['source_cells'].append(cell.coordinate)
                calendar_weeks.append(record)
                cells = [sheet.cell(r, col) for r in practice_rows
                         if isinstance(sheet.cell(r, col).value, (int, float)) and sheet.cell(r, col).value > 0]
                if cells:
                    practice_weeks.append({'from': first.isoformat(), 'to': last.isoformat(),
                                           'hours': sum(c.value for c in cells),
                                           'cells': [c.coordinate for c in cells],
                                           'fills': [str(c.fill.fgColor.index) for c in cells]})
            periods = []
            for week in practice_weeks:
                if periods and periods[-1]['to'] == (date.fromisoformat(week['from']) - timedelta(days=1)).isoformat():
                    periods[-1]['to'] = week['to']
                    periods[-1]['calendar_hours'] += week['hours']
                    periods[-1]['source_cells'] += week['cells']
                else:
                    periods.append({'from': week['from'], 'to': week['to'], 'calendar_hours': week['hours'],
                                    'source_cells': list(week['cells']), 'kind': 'industrial'})
            entries.append({'group_id': mapped['id'], 'group_name': mapped['name'], 'source_name': name,
                            'sheet': sheet.title, 'row': row, 'practice_periods': periods,
                            'calendar_theory_semester_hours': sheet.cell(row, 7).value,
                            'academic_calendar': calendar_weeks,
                            'practice_weeks': practice_weeks})
    return {'source_url': 'https://docs.google.com/spreadsheets/d/1YQ4TsERPrgtNNrC1weyDPox8hwIjU8XDfaM9XDO2X4U/edit',
            'source_sha256': hashlib.sha256(Path(source).read_bytes()).hexdigest(),
            'groups': entries, 'unmatched': unmatched, 'duplicates': duplicates,
            'missing_groups': [g['name'] for g in data['groups'] if g['id'] not in used]}


def apply(data, report, allow_partial=False):
    if not allow_partial and (report['unmatched'] or report['duplicates'] or report['missing_groups']):
        raise ValueError('Импорт остановлен: не все группы сопоставлены однозначно')
    result = copy.deepcopy(data)
    entries = {e['group_id']: e for e in report['groups']}
    for group in result['groups']:
        if group['id'] not in entries:
            continue
        entry = entries[group['id']]
        group['practice_periods'] = entry['practice_periods']
        group['academic_calendar'] = entry['academic_calendar']
        group['calendar_theory_semester_hours'] = entry['calendar_theory_semester_hours']
        group['practice_calendar_source'] = {'url': report['source_url'], 'sheet': entry['sheet'],
                                            'row': entry['row'], 'name': entry['source_name'],
                                            'sha256': report['source_sha256']}
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('source')
    parser.add_argument('--data', default='data/timetable_data.json')
    parser.add_argument('--report', required=True)
    parser.add_argument('--output')
    parser.add_argument('--confirm-aliases', action='store_true')
    parser.add_argument('--allow-partial', action='store_true')
    args = parser.parse_args()
    data = json.loads(Path(args.data).read_text(encoding='utf-8-sig'))
    report = extract(args.source, data, args.confirm_aliases)
    Path(args.report).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    if args.output:
        Path(args.output).write_text(json.dumps(apply(data, report, args.allow_partial), ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: v for k, v in report.items() if k != 'groups'}, ensure_ascii=False))
    for g in report['groups']:
        print(g['group_name'], [(p['from'], p['to']) for p in g['practice_periods']])
