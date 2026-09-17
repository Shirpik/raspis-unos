"""Read-only comparison of workload source and stable database identities."""
import json, sys
from pathlib import Path
from collections import defaultdict
import openpyxl
from merge_vkleyki_workload import clean, norm, numeric, columns, subgroup, subject_name

root = Path(__file__).resolve().parents[1]
source = Path(sys.argv[1])
out = root / 'outputs/workload-import-20260914-v5'
out.mkdir(parents=True, exist_ok=True)
data = json.loads((root/'data/timetable_data.json').read_text(encoding='utf-8-sig'))
groups = {norm(g['name']): g for g in data['groups']}
teachers = {t['id']: t['name'] for t in data['teachers']}
old = defaultdict(list)
for l in data['lessons']:
    old[(l['group'],norm(l['name']),l.get('subgroup',-1))].append(l)
rows=[]; missing=[]; ambiguous=[]; changed=[]; used=set(); rawteachers=set()
w = openpyxl.load_workbook(source,data_only=True)
for ws in w:
    g=groups[norm(ws.title)]
    h,s,t,c=columns(ws)
    assert min(h,s,t)>0
    sem=[(r,col,clean(ws.cell(r,col).value)) for r in range(1,13) for col in range(1,ws.max_column+1) if 'сем' in norm(ws.cell(r,col).value)]
    assert any(col==c and any(str(n) in v for n in (1,3,5,7)) for _,col,v in sem), (ws.title,sem)
    for r in range(h+1,ws.max_row+1):
        name=clean(ws.cell(r,s).value); hours=numeric(ws.cell(r,c).value)
        if not name or hours is None or hours<=0: continue
        teacher=clean(ws.cell(r,t).value); rawteachers.add(teacher)
        row=dict(sheet=ws.title,row=r,group=g['id'],name=subject_name(name),subgroup=subgroup(name,g['id']),hours=hours,teacher=teacher,index=clean(ws.cell(r,2).value))
        candidates=old[(row['group'],norm(row['name']),row['subgroup'])]
        exact=[l for l in candidates if norm(teachers.get(l.get('teacher'),''))==norm(teacher)]
        active=[l for l in candidates if l.get('curriculum_active',True)]
        candidates=exact or active or candidates
        if len(candidates)==1:
            l=candidates[0]; row['id']=l['id']; used.add(l['id'])
            dif={k:[a,b] for k,a,b in [('hours',l.get('total_hours'),hours),('teacher',teachers.get(l.get('teacher'),''),teacher)] if a!=b}
            if dif: changed.append(dict(**row,changes=dif))
        elif candidates: ambiguous.append(dict(**row,candidates=candidates))
        else: missing.append(row)
        rows.append(row)
unmatched=[dict(l,group_name=next(g['name'] for g in data['groups'] if g['id']==l['group']),teacher_name=teachers.get(l.get('teacher'),'')) for l in data['lessons'] if l['id'] not in used]
report=dict(rows=rows,changed=changed,new=missing,ambiguous=ambiguous,unmatched=unmatched,new_teachers=sorted(rawteachers-{t['name'] for t in data['teachers']}))
(out/'comparison.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(dict(sheets=len(w.sheetnames),rows=len(rows),changed=changed,new=missing,ambiguous=ambiguous,new_teachers=report['new_teachers'],unmatched=[{k:l.get(k) for k in ('id','name','group_name','teacher_name','total_hours','curriculum_active','plan_active')} for l in unmatched]),ensure_ascii=False,indent=2))
