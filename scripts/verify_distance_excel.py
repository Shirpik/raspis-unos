"""Read back final XLSX and reconcile every timetable cell to audited events."""
import json
import sys
from collections import Counter
from pathlib import Path
import openpyxl

p=Path(sys.argv[1] if len(sys.argv)>1 else 'outputs/solver-audit-20260918')
data=json.loads((p/'base-data.json').read_text(encoding='utf-8'))
schedule=json.loads((p/'candidate-cp/schedule_all.json').read_text(encoding='utf-8'))
teachers={t['id']:t for t in data['teachers']}
rooms={r['id']:r for r in data['rooms']}
groups={g['group_name']:g for g in schedule['groups']}
file=next(p.glob('*.xlsx'))
file=p/'Расписание_18-19.09.2026_дистант_без_дублей.xlsx'
wb=openpyxl.load_workbook(file,data_only=False)
source=openpyxl.load_workbook('frontend/public/templates/schedule-template.xlsx',data_only=False)
assert wb.sheetnames==source.sheetnames
totals=Counter(); inspected=set()
for sh in wb:
    for day,header in [('2026-09-18',62),('2026-09-19',77)]:
        assert sh.cell(header,1).value.strftime('%Y-%m-%d')==day
        headers={c:str(sh.cell(header,c).value or '').strip() for c in range(4,16)}
        headers={c:n for c,n in headers.items() if n}
        expected={}
        for c,name in headers.items():
            group=groups[name]; inspected.add(name)
            d=next(x for x in group['days'] if x['date_iso']==day)
            for slot in d['slots']:
                for l in slot['lessons']:
                    ordinal=0 if l['subgroup']<0 else l['subgroup']%2+1
                    row=header+1+2*(slot['slot']-1)+(ordinal==2)
                    room=rooms.get(l['room_id'])
                    place=f"{room['name']}_{'Л' if room['campus']==0 else 'К'}" if room and data['settings'].get('distance_revision',1)>=2 else 'Дистант'
                    text=l['name']+(f' {ordinal} п/г' if ordinal else '')+'\n'+teachers[l['teacher_id']]['name'].split()[0]+' '+place
                    assert (row,c) not in expected
                    expected[row,c]=text
                    if ordinal==0: assert sh.cell(row+1,c).value is None, (sh.title,row,c,'duplicate common lesson')
        for row in range(header+1,header+15):
            for c in headers:
                actual=sh.cell(row,c).value
                assert actual==expected.get((row,c)),(sh.title,row,c,actual,expected.get((row,c)))
                if actual: totals[day]+=1
        assert not any(r.min_row<=header+14 and r.max_row>=header+1 and r.max_col>=4 for r in sh.merged_cells.ranges)
    # Values and formulas outside the requested Friday/Saturday block remain.
    for row in source[sh.title]:
        for cell in row:
            if 62<=cell.row<=91: continue
            assert sh[cell.coordinate].value==cell.value,(sh.title,cell.coordinate,'unrelated content changed')
assert inspected==set(groups)
audit=json.loads((p/'final-validation.json').read_text(encoding='utf-8'))
assert audit['ok'] and dict(totals)==audit['pairs_by_date'] and sum(totals.values())==audit['total_pairs']
if data['settings'].get('distance_revision',1)>=2:
    for sh in wb:
        for row in (62,77):
            assert sh.cell(row,4).fill.fgColor.rgb=='FFA9D18E'
            assert sh.cell(row+1,4).fill.fgColor.rgb=='FFE2F0D9'
report={'ok':True,'pairs_by_date':dict(totals),'total_pairs':sum(totals.values()),'groups':len(inspected),'all_cells_match':True,'common_lower_cells_empty':True,'unrelated_values_preserved':True}
(p/'excel-validation.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(report)
